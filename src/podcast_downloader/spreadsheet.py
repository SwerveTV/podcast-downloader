from __future__ import annotations

import csv
import io
import re
import urllib.parse
import urllib.request
from collections.abc import Iterable
from pathlib import Path

from openpyxl import load_workbook  # type: ignore[import-untyped]

from .exceptions import InvalidSpreadsheetError
from .models import ShowJob

CANONICAL_COLUMNS = {
    "network / distributor": "network",
    "network": "network",
    "distributor": "network",
    "show name": "show_name",
    "show": "show_name",
    "host(s)": "hosts",
    "hosts": "hosts",
    "number of episodes": "number_of_episodes",
    "subscriber count": "subscriber_count",
    "full episode playlist": "playlist_url",
    "playlist": "playlist_url",
    "playlist url": "playlist_url",
    "operator": "operator",
    "status": "status",
}


def normalize_header(value: object) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"\s+", " ", text.strip()).strip("\ufeff")


def _canonical_key(header: str) -> str | None:
    cleaned = normalize_header(header).lower()
    if not cleaned or cleaned.startswith("unnamed:"):
        return None
    return CANONICAL_COLUMNS.get(cleaned)


def _row_values(row: Iterable[object]) -> list[str]:
    return ["" if cell is None else str(cell).strip() for cell in row]


def read_spreadsheet(source: str | Path) -> tuple[list[ShowJob], list[str], list[str]]:
    """Return jobs, ignored columns, and raw normalized headers."""

    rows = _load_rows(source)
    if not rows:
        raise InvalidSpreadsheetError("Spreadsheet is empty.")

    header_index = _find_header_row(rows)
    if header_index is None:
        raise InvalidSpreadsheetError("Could not find a header row with Show Name and Full Episode Playlist.")

    headers = [normalize_header(v) for v in rows[header_index]]
    mapped: dict[int, str] = {}
    ignored: list[str] = []
    for idx, header in enumerate(headers):
        key = _canonical_key(header)
        if key:
            mapped[idx] = key
        elif header and not header.lower().startswith("unnamed:"):
            ignored.append(header)

    if "show_name" not in mapped.values() or "playlist_url" not in mapped.values():
        raise InvalidSpreadsheetError("Spreadsheet must include Show Name and Full Episode Playlist columns.")

    jobs: list[ShowJob] = []
    current_network = ""
    for row_number, raw_row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        values = _row_values(raw_row)
        if not any(values):
            continue
        data = {key: "" for key in CANONICAL_COLUMNS.values()}
        raw: dict[str, str] = {}
        for idx, value in enumerate(values):
            header = headers[idx] if idx < len(headers) else ""
            if idx in mapped:
                data[mapped[idx]] = value.strip()
            elif header and value.strip():
                raw[header] = value.strip()

        if data["network"]:
            current_network = data["network"]
        elif current_network:
            data["network"] = current_network

        is_section = _looks_like_section_header(data, values, mapped)
        jobs.append(
            ShowJob(
                row_number=row_number,
                network=data["network"],
                show_name=data["show_name"],
                hosts=data["hosts"],
                number_of_episodes=data["number_of_episodes"],
                subscriber_count=data["subscriber_count"],
                playlist_url=data["playlist_url"],
                operator=data["operator"],
                status=data["status"],
                raw=raw,
                is_section_header=is_section,
            )
        )
    return jobs, sorted(set(ignored)), headers


def _looks_like_section_header(data: dict[str, str], values: list[str], mapped: dict[int, str]) -> bool:
    if data["playlist_url"].strip():
        return False
    meaningful = [v for v in values if v.strip()]
    if len(meaningful) > 2:
        return False
    show_idx = next((idx for idx, key in mapped.items() if key == "show_name"), None)
    network_idx = next((idx for idx, key in mapped.items() if key == "network"), None)
    if network_idx is not None and network_idx < len(values) and values[network_idx].strip():
        return True
    return show_idx is not None and show_idx < len(values) and bool(values[show_idx].strip())


def _find_header_row(rows: list[list[object]]) -> int | None:
    for idx, row in enumerate(rows[:25]):
        keys = {_canonical_key(normalize_header(cell)) for cell in row}
        if "show_name" in keys and "playlist_url" in keys:
            return idx
    return None


def _load_rows(source: str | Path) -> list[list[object]]:
    source_text = str(source)
    if source_text.startswith("http://") or source_text.startswith("https://"):
        with urllib.request.urlopen(_google_export_url(source_text), timeout=30) as response:
            text = response.read().decode("utf-8-sig")
        return [list(row) for row in csv.reader(io.StringIO(text))]

    path = Path(source).expanduser()
    if not path.exists():
        raise InvalidSpreadsheetError(f"Input file does not exist: {path}")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return [list(row) for row in csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig")))]
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    raise InvalidSpreadsheetError("Input must be a .csv, .xlsx, or public Google Sheets URL.")


def _google_export_url(url: str) -> str:
    parsed = urllib.parse.urlparse(url)
    if "docs.google.com" not in parsed.netloc:
        return url
    match = re.search(r"/spreadsheets/d/([^/]+)", parsed.path)
    if not match:
        raise InvalidSpreadsheetError("Google Sheets URL does not contain a spreadsheet ID.")
    query = urllib.parse.parse_qs(parsed.query)
    gid = query.get("gid", ["0"])[0]
    return f"https://docs.google.com/spreadsheets/d/{match.group(1)}/export?format=csv&gid={gid}"
